"""Read a State Gas board deck into structured JSON: every slide's title,
paragraph bullets (with indent level) and tables. Used both to build the
style bank and to give each month's run last month's deck as input."""
import re
from pptx import Presentation


def _title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if re.match(r"^(\d+(\.\d+)?|[A-Z]|Appendix)\.?\s*\t", t) or re.match(r"^\d+(\.\d+)?\.\s", t):
                return re.sub(r"\s+", " ", t.split("\n")[0]).strip()
    return ""


def read_deck(path):
    p = Presentation(path)
    slides = []
    for i, s in enumerate(p.slides, 1):
        title = _title(s)
        texts, tables = [], []
        for sh in sorted(s.shapes, key=lambda x: ((x.top or 0) // 300000, x.left or 0)):
            if sh.has_table:
                grid = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                tables.append(grid)
            elif sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if not t or re.sub(r"\s+", " ", t.split("\n")[0]).strip() == title:
                    continue
                paras = [{"level": pa.level, "text": "".join(r.text for r in pa.runs).strip()}
                         for pa in sh.text_frame.paragraphs if "".join(r.text for r in pa.runs).strip()]
                bold = any(r.font.bold for pa in sh.text_frame.paragraphs for r in pa.runs)
                texts.append({"box": sh.name, "heading": bold and len(paras) == 1 and len(paras[0]["text"]) < 80,
                              "x": round((sh.left or 0) / 914400, 2), "y": round((sh.top or 0) / 914400, 2),
                              "paras": paras})
        slides.append({"n": i, "title": title, "texts": texts, "tables": tables})
    return slides


def deck_as_text(slides, max_chars=40000):
    out = []
    for s in slides:
        out.append(f"=== SLIDE {s['n']}: {s['title']}")
        for t in s["texts"]:
            for p in t["paras"]:
                out.append(("    " * p["level"]) + ("## " if t["heading"] else "- ") + p["text"])
        for g in s["tables"]:
            for r in g:
                out.append(" | ".join(c.replace("\n", " / ") for c in r))
    return "\n".join(out)[:max_chars]
