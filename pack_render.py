"""
pack_render.py - draw the monthly board deck in the MD's own template.

The template is his latest final deck: its single slide layout carries the
State Gas logo, crimson footer and confidentiality line. We keep that layout,
drop the old slides, and draw each page with his measurements and type:
Georgia Pro Light, 14pt bold titles with a short crimson rule, 8-9pt body.
Long pages and tables continue onto "N.1 (continued)" pages, as he does.
"""
import copy
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FONT = "Georgia Pro Light"
CRIMSON = RGBColor(0xC8, 0x29, 0x3F)
GREY_LINE = RGBColor(0xE8, 0xE8, 0xE8)
GREY_FILL = RGBColor(0xF2, 0xF2, 0xF2)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x80, 0x80, 0x80)
BODY_TOP, BODY_BOTTOM = 0.47, 5.30          # usable band between title and footer
CH_PER_IN = {7: 24.0, 8: 21.0, 9: 19.0, 10: 16.5}   # rough Georgia Pro Light widths
LINE_IN = {7: 0.12, 8: 0.135, 9: 0.15, 10: 0.17}


def _lines(text, width_in, pt):
    per = max(8, int(width_in * CH_PER_IN[pt]))
    return sum(max(1, math.ceil(len(part) / per)) for part in (text or " ").split("\n"))


def text_height(paras, width_in, pt):
    return sum(_lines(("    " * p.get("level", 0)) + p["text"], width_in - 0.25 * p.get("level", 0), pt)
               for p in paras) * LINE_IN[pt] + 0.04 * len(paras) + 0.1


class Deck:
    def __init__(self, template_path):
        self.prs = Presentation(template_path)
        self.logo = None
        first = self.prs.slides[0] if len(self.prs.slides) else None
        if first:
            for sh in first.shapes:
                if sh.shape_type == 13:              # the big logo on his title page
                    self.logo = (sh.image.blob, sh.left, sh.top, sh.width, sh.height)
                    break
        # drop every existing slide, keep the layout (logo, footer, CONFIDENTIAL)
        sldIdLst = self.prs.slides._sldIdLst
        for sldId in list(sldIdLst):
            self.prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)
        self.layout = self.prs.slide_layouts[0]

    # ── primitives ──────────────────────────────────────────────────────────
    def _slide(self):
        return self.prs.slides.add_slide(self.layout)

    def _box(self, s, x, y, w, h, paras, pt=8, bold=False, color=INK, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right"):
            setattr(tf, m, Inches(0.05))
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            lvl = p.get("level", 0)
            para.level = min(lvl, 4)
            para.space_after = Pt(3)
            if p.get("bullet", True) and not bold:
                self._bullet(para, lvl)
            r = para.add_run()
            r.text = p["text"]
            f = r.font
            f.name, f.size, f.bold = FONT, Pt(pt), bold or p.get("bold", False)
            f.color.rgb = color
        return tb

    @staticmethod
    def _bullet(para, lvl):
        from pptx.oxml.ns import qn
        pPr = para._p.get_or_add_pPr()
        indent = 0.17 + 0.2 * lvl
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", str(int(-Inches(0.15))))
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
            for e in pPr.findall(qn(tag)):
                pPr.remove(e)
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•" if lvl == 0 else "–"})
        pPr.append(bu)

    def _title(self, s, number, title):
        label = f"{number}." if number and "." not in str(number) else str(number or "")
        self._box(s, 0.40, 0.08, 8.2, 0.26, [{"text": f"{label}\t{title}" if label else title, "bullet": False}],
                  pt=14, bold=True)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.40), Inches(0.37), Inches(0.51), Inches(0.03))
        bar.fill.solid(); bar.fill.fore_color.rgb = CRIMSON; bar.line.fill.background()

    def _placeholder(self, s, x, y, w, h, note):
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid(); r.fill.fore_color.rgb = GREY_FILL
        r.line.color.rgb = MUTED; r.line.dash_style = 4
        tf = r.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = f"Insert: {note}"
        run.font.name, run.font.size, run.font.italic = FONT, Pt(9), True
        run.font.color.rgb = MUTED

    def _table(self, s, x, y, w, header, rows, widths, pt=8, first_col_bold=False):
        shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(0.3))
        tbl = shape.table
        tbl.first_row = True
        # plain style like his: no banding, thin grey grid
        tblPr = tbl._tbl.tblPr
        tblPr.set("bandRow", "0")
        for i, wd in enumerate(widths):
            tbl.columns[i].width = Inches(wd)
        for ri, row in enumerate([header] + rows):
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.fill.background()
                cell.margin_left = cell.margin_right = Inches(0.04)
                cell.margin_top = cell.margin_bottom = Inches(0.02)
                tf = cell.text_frame; tf.word_wrap = True
                parts = [t for t in str(val or "").split("\n")] or [""]
                for pi, part in enumerate(parts):
                    para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                    r = para.add_run(); r.text = part
                    r.font.name, r.font.size = FONT, Pt(pt)
                    r.font.bold = ri == 0 or (first_col_bold and ci == 0)
                    r.font.color.rgb = INK
        self._grid(tbl)
        return tbl

    @staticmethod
    def _grid(tbl):
        from pptx.oxml.ns import qn
        for r in tbl.rows:
            for c in r.cells:
                tcPr = c._tc.get_or_add_tcPr()
                for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                    ln = tcPr.makeelement(qn(side), {"w": "6350"})
                    sf = ln.makeelement(qn("a:solidFill"), {})
                    clr = sf.makeelement(qn("a:srgbClr"), {"val": "BFBFBF"})
                    sf.append(clr); ln.append(sf); tcPr.append(ln)

    # ── pages ───────────────────────────────────────────────────────────────
    def title_page(self, meeting_date, author):
        s = self._slide()
        s._element.set("showMasterSp", "0")          # his title page: big logo only
        band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(5.36), Inches(10), Inches(0.27))
        band.fill.solid(); band.fill.fore_color.rgb = RGBColor(0xA5, 0x1C, 0x30); band.line.fill.background()
        self._box(s, 0.41, 5.36, 9.47, 0.26, [{"text": "CONFIDENTIAL – FOR BOARD USE ONLY", "bullet": False}],
                  pt=8, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), anchor=MSO_ANCHOR.MIDDLE)
        self._box(s, 0.41, 2.13, 5.94, 1.10, [{"text": "Executive Board Report", "bullet": False}], pt=10, bold=True)
        s.shapes[-1].text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        self._box(s, 0.47, 4.03, 7.0, 0.45, [{"text": meeting_date, "bullet": False}], pt=10)
        s.shapes[-1].text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.47), Inches(4.58), Inches(3.5), Inches(0.04))
        bar.fill.solid(); bar.fill.fore_color.rgb = CRIMSON; bar.line.fill.background()
        self._box(s, 0.47, 4.83, 5.0, 0.35, [{"text": f"Author: {author}", "bullet": False}], pt=10)
        if self.logo:
            import io
            blob, x, y, w, h = self.logo
            s.shapes.add_picture(io.BytesIO(blob), x, y, w, h)

    def exec_summary(self, number, headlines, boxes):
        s = self._slide()
        self._title(s, number, "Executive Summary")
        hp = [{"text": t} for t in headlines]
        hh = min(1.6, text_height(hp, 9.28, 8))
        self._box(s, 0.40, 0.47, 9.28, hh, hp, pt=8)
        top = 0.47 + hh + 0.06
        n = max(1, len(boxes))
        cols = 1 if n == 1 else 2
        rows = math.ceil(n / cols)
        gap = 0.10
        bw = (9.44 - gap * (cols - 1)) / cols
        bh = (BODY_BOTTOM - top - gap * (rows - 1)) / rows
        for i, b in enumerate(boxes):
            r, c = divmod(i, cols)
            if i == n - 1 and n % 2 == 1 and cols == 2:   # odd last box spans the row
                c, w = 0, 9.44
            else:
                w = bw
            x, y = 0.31 + c * (bw + gap), top + r * (bh + gap)
            frame = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(bh))
            frame.fill.solid(); frame.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            frame.line.color.rgb = GREY_LINE
            head = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.26))
            head.fill.solid(); head.fill.fore_color.rgb = GREY_LINE; head.line.fill.background()
            self._box(s, x + 0.08, y, w - 0.16, 0.26, [{"text": b["title"], "bullet": False}],
                      pt=10, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            self._box(s, x + 0.05, y + 0.30, w - 0.10, bh - 0.32, [{"text": t} for t in b["bullets"]], pt=8)
        return s

    def op_plan(self, number, title, columns, rows):
        s = self._slide()
        self._title(s, number, title)
        lw = 0.62
        cw = (9.31 - lw) / len(columns)
        self._table(s, 0.34, 0.48, 9.31, [""] + columns, [[r["label"]] + r["cells"] for r in rows],
                    [lw] + [cw] * len(columns), pt=7, first_col_bold=True)

    def narrative(self, number, title, paras, table=None, image_note=None, intro=None):
        """A strategy page. Splits onto continuation pages when long."""
        pages, cur, used = [], [], 0.0
        width = 4.9 if image_note else 9.23
        cap = BODY_BOTTOM - BODY_TOP - 0.1 - (0.3 if intro else 0)
        for p in paras:
            h = text_height([p], width, 9)
            if cur and used + h > cap:
                pages.append(cur); cur, used = [], 0.0
            cur.append(p); used += h
        if cur or not pages:
            pages.append(cur)
        for i, pp in enumerate(pages):
            s = self._slide()
            t = title if i == 0 else f"{title} (continued)"
            num = number if i == 0 else f"{number}.{i}"
            self._title(s, num, t)
            y = BODY_TOP + 0.08
            if intro and i == 0:
                self._box(s, 0.40, y, 9.23, 0.3, [{"text": intro, "bullet": False}], pt=9); y += 0.32
            if image_note and i == 0:
                self._placeholder(s, 0.40, y, 4.1, BODY_BOTTOM - y - 0.05, image_note)
                self._box(s, 4.65, y, 5.0, BODY_BOTTOM - y, pp, pt=9)
            else:
                self._box(s, 0.40, y, 9.23, BODY_BOTTOM - y, pp, pt=9)
        if table:
            self.table_pages(f"{number}.{len(pages)}", f"{title} (continued)", table["header"], table["rows"],
                             table.get("widths"), intro=table.get("intro"))

    def table_pages(self, number, title, header, rows, widths=None, intro=None, pt=8):
        total_w = 9.23
        if not widths:
            widths = [total_w / len(header)] * len(header)
        cap = BODY_BOTTOM - BODY_TOP - 0.15 - (0.32 if intro else 0)
        def rh(row):
            return max(_lines(str(v), w - 0.08, pt) for v, w in zip(row, widths)) * LINE_IN[pt] + 0.06
        head_h = rh(header)
        chunks, cur, used = [], [], head_h
        for r in rows:
            h = rh(r)
            if cur and used + h > cap:
                chunks.append(cur); cur, used = [], head_h
            cur.append(r); used += h
        chunks.append(cur)
        for i, ch in enumerate(chunks):
            s = self._slide()
            self._title(s, number, title)
            y = BODY_TOP + 0.02
            if intro and i == 0:
                self._box(s, 0.40, y, total_w, 0.3, [{"text": intro, "bullet": False}], pt=9); y += 0.32
            self._table(s, 0.40, y, total_w, header, ch, widths, pt=pt)

    def cashflow(self, number, title, bullets, image_note):
        s = self._slide()
        self._title(s, number, title)
        self._placeholder(s, 0.40, 0.50, 6.55, 4.66, image_note)
        self._box(s, 7.05, 0.47, 2.83, 4.83, [{"text": b} for b in bullets], pt=8)

    def save(self, path):
        self.prs.save(path)


def build(spec, template, out_path):
    d = Deck(template)
    d.title_page(spec["meeting_date"], spec.get("author", "Doug McAlpine"))
    n = 1
    d.exec_summary(n, spec["exec"]["headlines"], spec["exec"]["boxes"]); n += 1
    op = spec.get("op_plan")
    if op:
        d.op_plan(n, op["title"], op["columns"], op["rows"]); n += 1
    for pg in spec.get("strategy_pages", []):
        d.narrative(n, pg["title"], pg["paras"], pg.get("table"), pg.get("image_note"), pg.get("intro")); n += 1
    cf = spec.get("cashflow")
    if cf:
        d.cashflow(n, cf["title"], cf["bullets"], cf.get("image_note", "cashflow forecast chart from the model")); n += 1
    om = spec.get("other_matters")
    if om:
        d.table_pages(n, "Other matters", ["Area", "Matter", "Update", "Next Steps", "Risk", "Action"],
                      [[r.get("area", ""), r.get("matter", ""), r.get("update", ""), r.get("next_steps", ""),
                        r.get("risk", ""), r.get("action", "")] for r in om],
                      [0.55, 0.85, 3.95, 2.78, 0.55, 0.55])
    d.save(out_path)
    return out_path
